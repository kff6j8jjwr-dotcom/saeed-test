from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

output_path = "assets/eobi-information-request-letter-template.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.95)
section.right_margin = Inches(0.95)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(9)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run("EOBI Information Request Letter")
title_run.font.name = "Aptos Display"
title_run.font.color.rgb = None

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.add_run("Editable template")
subtitle_run.italic = True
subtitle_run.font.size = Pt(10)

for line in ("[Your name]", "[Your postal address]", "[City]", "[Email address] | [Phone number]", "[Date]"):
    doc.add_paragraph(line)

doc.add_paragraph("To\nThe Chairman / Concerned Officer\nEmployees' Old-Age Benefits Institution (EOBI)\n[Office address]")

subject = doc.add_paragraph()
subject_run = subject.add_run("Subject: Request for information regarding EOBI registration, contributions and benefits")
subject_run.bold = True

doc.add_paragraph("Dear Sir/Madam,")

doc.add_paragraph(
    "I am writing to request clear information concerning my EOBI record and the process for accessing the benefits for which I may be eligible. Please treat this letter as a request for guidance and, where applicable, copies or details of the relevant record."
)

doc.add_paragraph("Kindly provide the following information:")
for item in (
    "My registration status and EOBI number, if available.",
    "The contribution record submitted by my employer(s), including the periods covered.",
    "My current eligibility for old-age pension, survivors' pension, invalidity pension or old-age grant, as applicable.",
    "Any missing information, documents or corrective steps required to complete or update my record.",
    "The appropriate office, procedure and expected timeframe for resolving this matter.",
):
    doc.add_paragraph(item, style="List Bullet")

doc.add_paragraph(
    "My relevant details are: CNIC [number], employer [name], employment period [dates], and EOBI number [if known]. I would appreciate a written response at the contact details above."
)

doc.add_paragraph("Yours faithfully,")
doc.add_paragraph("\n[Signature]\n[Your name]")

doc.core_properties.title = "EOBI Information Request Letter"
doc.core_properties.author = "Saeed Rasheed"
doc.save(output_path)
